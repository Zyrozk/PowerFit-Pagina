from datetime import date, datetime, time, timedelta

from django.utils import timezone
from rest_framework import generics
from rest_framework.exceptions import ValidationError
from rest_framework.permissions import BasePermission, IsAuthenticated

from .models import Pago
from .serializers import PagoSerializer

from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


class EsDuena(BasePermission):
    def has_permission(self, request, view):
        return (
            request.user.is_authenticated
            and request.user.rol == "DUENA"
        )


class HistorialPagosView(generics.ListAPIView):
    serializer_class = PagoSerializer
    permission_classes = [IsAuthenticated, EsDuena]

    def get_queryset(self):
        pagos = Pago.objects.select_related("socio").order_by("-fecha_pago")

        fecha_inicio = self.request.query_params.get("fecha_inicio")
        fecha_fin = self.request.query_params.get("fecha_fin")

        def convertir_fecha(valor):
            try:
                return date.fromisoformat(valor)
            except ValueError:
                raise ValidationError({
                    "fecha": "La fecha debe tener el formato AAAA-MM-DD."
                })

        if fecha_inicio:
            inicio = convertir_fecha(fecha_inicio)
        else:
            inicio = None

        if fecha_fin:
            fin = convertir_fecha(fecha_fin)
        else:
            fin = None

        if inicio and fin and inicio > fin:
            raise ValidationError({
                "fecha": "La fecha de inicio no puede ser posterior a la fecha de término."
            })

        if inicio:
            inicio_datetime = timezone.make_aware(
                datetime.combine(inicio, time.min),
                timezone.get_current_timezone()
            )
            pagos = pagos.filter(fecha_pago__gte=inicio_datetime)

        if fin:
            fin_exclusivo = timezone.make_aware(
                datetime.combine(fin + timedelta(days=1), time.min),
                timezone.get_current_timezone()
            )
            pagos = pagos.filter(fecha_pago__lt=fin_exclusivo)

        return pagos


class ExportarPagosExcelView(HistorialPagosView):

    def get(self, request, *args, **kwargs):
        pagos = self.get_queryset()

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Historial de pagos"

        encabezados = [
            "ID",
            "Socio",
            "Tipo de pago",
            "Método de pago",
            "Monto (CLP)",
            "Estado",
            "Fecha de pago",
        ]

        hoja.append(encabezados)

        for celda in hoja[1]:
            celda.font = Font(color="FFFFFF", bold=True)
            celda.fill = PatternFill(
                start_color="198754",
                end_color="198754",
                fill_type="solid"
            )

        for pago in pagos:
            fecha_local = timezone.localtime(pago.fecha_pago)

            hoja.append([
                pago.id,
                pago.socio.get_full_name() or pago.socio.username,
                pago.get_tipo_pago_display(),
                pago.get_metodo_pago_display(),
                pago.monto,
                pago.get_estado_display(),
                fecha_local.strftime("%d-%m-%Y %H:%M"),
            ])

        anchos = {
            "A": 10,
            "B": 30,
            "C": 22,
            "D": 25,
            "E": 18,
            "F": 16,
            "G": 22,
        }

        for columna, ancho in anchos.items():
            hoja.column_dimensions[columna].width = ancho

        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions

        for fila in hoja.iter_rows(min_row=2):
            fila[4].number_format = '"$"#,##0'

        archivo = BytesIO()
        libro.save(archivo)
        archivo.seek(0)

        respuesta = HttpResponse(
            archivo.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

        respuesta["Content-Disposition"] = (
            'attachment; filename="historial_pagos_powerfit.xlsx"'
        )

        return respuesta 